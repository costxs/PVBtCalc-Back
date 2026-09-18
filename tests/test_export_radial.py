"""
tests/test_export_radial.py

Export radial com figuras matplotlib (PNG 300 dpi) embutidas no workbook e
no pacote de figuras (.zip). Caso de teste (pedido explicito):
    Indiana Limestone | HCl With Inhibitor Corrosion | 338.71 K
    alvos 5/10/15/20 ft | sweep 0.0024-0.1 bbl/min | 50 passos
    temperaturas de comparacao 297.04/338.71/422.04 K | vazoes skin 0.8/1.6/3.2

REGRA verificada aqui tambem: export_plots.py/export_workbook.py nunca
importam PVBTradialFunc -- os dados vem de calculate_pvbt_radial/
generate_design_plot/generate_skin_evolution (as MESMAS rotinas que o
endpoint HTTP usa), e o modulo de export so redesenha o que ja saiu delas.

    venv/Scripts/python.exe -m pytest tests/test_export_radial.py -v
"""
import io
import os
import sys
import xml.etree.ElementTree as ET
import zipfile

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pytest
from PIL import Image

from app.routes.pvbtRadialCurve import calculate_pvbt_radial
from app.routes.designPlot import generate_design_plot
from app.services.tools import generate_skin_evolution
from app.services import export_plots as ep
from app.services import export_workbook as ewb
from app.schemas import (
    DesignPlotInput,
    DesignPlotSeries,
    ExportInputs,
    RadialCurveInput,
    RadialCurveResult,
    RadialExportOptions,
    RadialExportRequest,
    RadialGeometryInput,
    RadialSystem,
    RadialTargetsInput,
    FlowrateSweepInput,
    SkinEvolutionInput,
    SkinEvolutionPoint,
)

ROCK = "Indiana Limestone"
ACID = "HCl With Inhibitor Corrosion"
CONC = 0.15
POROSITY = 0.15
TEMP_K = 338.71
TARGETS = [5.0, 10.0, 15.0, 20.0]
SWEEP_MIN_BBL = 0.0024
SWEEP_MAX_BBL = 0.1
STEPS = 50
WELLBORE_RADIUS_IN = 1.5
PAYZONE_FT = 1.0
DESIGN_TEMPS = [297.04, 338.71, 422.04]
SKIN_FLOWS = [0.8, 1.6, 3.2]

MM_PER_INCH = 25.4
# bbox_inches="tight" recorta margem: aceita ate ~10% menor que o nominal,
# nunca maior (senao a figura nao caberia na coluna da revista).
SIZE_TOLERANCE = 0.10


def _system():
    return RadialSystem(rock_type=ROCK, porosity=POROSITY, acid_system=ACID,
                         acid_concentration=CONC, temperature_k=TEMP_K)


def _geometry():
    return RadialGeometryInput(wellbore_radius_in=WELLBORE_RADIUS_IN,
                                payzone_thickness_ft=PAYZONE_FT, drainage_radius_ft=None)


@pytest.fixture(scope="module")
def curve_output():
    data = RadialCurveInput(
        simulation_id="TEST1", system=_system(), geometry=_geometry(),
        radial_targets=RadialTargetsInput(target_mode="length", targets=TARGETS),
        flowrate_sweep=FlowrateSweepInput(min=SWEEP_MIN_BBL, max=SWEEP_MAX_BBL, steps=STEPS),
    )
    return calculate_pvbt_radial(data)


@pytest.fixture(scope="module")
def design_output():
    data = DesignPlotInput(
        simulation_id="TEST1", system=_system(), geometry=_geometry(),
        radial_targets=RadialTargetsInput(target_mode="length", targets=TARGETS),
        flowrate_sweep=FlowrateSweepInput(min=SWEEP_MIN_BBL, max=SWEEP_MAX_BBL, steps=STEPS),
        temperatures_to_compare=DESIGN_TEMPS,
    )
    return generate_design_plot(data)


@pytest.fixture(scope="module")
def skin_output():
    data = SkinEvolutionInput(
        acid_type=ACID, acid_concentration=CONC, temperature_k=TEMP_K,
        core_porosity=POROSITY, wellbore_radius_in=WELLBORE_RADIUS_IN,
        payzone_thickness_ft=PAYZONE_FT, rock_type=ROCK, flowrates_to_compare=SKIN_FLOWS,
    )
    return generate_skin_evolution(data, SKIN_FLOWS)


@pytest.fixture(scope="module")
def export_request(curve_output, design_output, skin_output) -> RadialExportRequest:
    curves = [RadialCurveResult(**c) if not isinstance(c, RadialCurveResult) else c
              for c in curve_output["curves"]]
    design_series = [DesignPlotSeries(**s) if not isinstance(s, DesignPlotSeries) else s
                      for s in design_output["series"]]
    skin_series = {
        key: [SkinEvolutionPoint(x=p["x"], y=p["y"], l_ft=p["l_ft"]) for p in points]
        for key, points in skin_output.items()
    }
    all_q = [q for c in curves for q in c.flowratepoints]

    return RadialExportRequest(
        inputs=ExportInputs(
            simulation_id="TEST1", rock_type=ROCK, acid_type=ACID,
            acid_concentration=CONC, porosity=POROSITY, temperature_k=TEMP_K,
            wellbore_size_in=3.0, wellbore_mode="diameter",
            wellbore_radius_in=WELLBORE_RADIUS_IN, payzone_thickness_ft=PAYZONE_FT,
            flowrate_min_bbl_min=SWEEP_MIN_BBL, flowrate_max_bbl_min=SWEEP_MAX_BBL,
            number_of_steps=STEPS,
            targets_label=", ".join(c.target_label for c in curves),
        ),
        curves=curves,
        design_series=design_series,
        skin_series=skin_series,
        options=RadialExportOptions(target_lengths=TARGETS, show_optimum_path=True, show_validity_band=True),
    )


# --- Workbook ----------------------------------------------------------

def test_workbook_sheet_order_and_names(export_request):
    import openpyxl
    xlsx_bytes = ewb.build_workbook(export_request)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

    expected = [
        "Inputs", "Resumo gráficos",
        "Design 297 K", "Design 339 K", "Design 422 K",
        "Sim 5.00 ft", "Sim 10.00 ft", "Sim 15.00 ft", "Sim 20.00 ft",
        "Skin 0.8 bbl-min", "Skin 1.6 bbl-min", "Skin 3.2 bbl-min",
    ]
    assert wb.sheetnames == expected


def test_workbook_freeze_panes_on_every_data_sheet(export_request):
    import openpyxl
    xlsx_bytes = ewb.build_workbook(export_request)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    for name in wb.sheetnames:
        if name == "Resumo gráficos":
            continue
        assert wb[name].freeze_panes == "A2", f"{name} sem freeze panes"


def test_workbook_inputs_are_numeric_not_text(export_request):
    import openpyxl
    xlsx_bytes = ewb.build_workbook(export_request)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Inputs"]
    numeric_labels = {
        "Acid Concentration (w/w)", "Porosity", "Temperature (K)", "Temperature (°C)",
        "Wellbore Size (in) [diameter]", "Wellbore Radius (in)", "Payzone Thickness (ft)",
        "Flowrate Sweep Min (bbl/min)", "Flowrate Sweep Max (bbl/min)",
        "Flowrate Sweep Min (gal/(ft.min))", "Flowrate Sweep Max (gal/(ft.min))",
        "Number of steps",
    }
    seen = set()
    for row in ws.iter_rows(min_row=2, max_col=2):
        label, value = row[0].value, row[1].value
        if label in numeric_labels:
            seen.add(label)
            assert isinstance(value, (int, float)), f"{label} veio como {type(value).__name__}: {value!r}"
    assert seen == numeric_labels


def test_workbook_one_image_per_data_sheet_plus_summary(export_request):
    xlsx_bytes = ewb.build_workbook(export_request)
    z = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    media = [n for n in z.namelist() if n.startswith("xl/media/")]
    # 10 abas de dados (3 design + 4 sim + 3 skin) + 3 combinadas no resumo
    assert len(media) == 10 + 3


DRAWING_NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
EMU_PER_PX = 9525  # unidade nativa do OOXML pra drawings; 9525 EMU = 1 px a 96 dpi


def test_workbook_image_display_width_and_summary_no_overlap(export_request):
    """Bug real (2026-09): XlsxWriter ja reduz a imagem pela razao 96/dpi_png
    antes de aplicar x_scale/y_scale -- calcular a escala sobre a largura em
    pixels do arquivo (sem contar o dpi=300 do figure pack) faz a imagem
    aparecer ~3x menor no Excel do que o alvo. Este teste le o XML de
    drawing (xl/drawings/*.xml) e confere a largura EXIBIDA (cx em EMU / 9525)
    -- nao so que a imagem foi inserida."""
    xlsx_bytes = ewb.build_workbook(export_request)
    z = zipfile.ZipFile(io.BytesIO(xlsx_bytes))

    data_sheet_widths_px = []
    resumo_boxes = []  # (y_emu, cy_emu, cx_emu) pra checar largura e sobreposicao

    for name in z.namelist():
        if not (name.startswith("xl/drawings/") and name.endswith(".xml")):
            continue
        root = ET.fromstring(z.read(name))
        for anchor in root.findall("xdr:twoCellAnchor", DRAWING_NS):
            from_col = int(anchor.find("xdr:from/xdr:col", DRAWING_NS).text)
            ext = anchor.find(".//a:ext", DRAWING_NS)
            off = anchor.find(".//a:off", DRAWING_NS)
            cx, cy = int(ext.attrib["cx"]), int(ext.attrib["cy"])
            y_emu = int(off.attrib["y"])

            if from_col == ewb.IMAGE_ANCHOR_COL:
                data_sheet_widths_px.append(cx / EMU_PER_PX)
            elif from_col == 0:
                resumo_boxes.append((y_emu, cy, cx))

    assert data_sheet_widths_px, "nenhuma imagem encontrada nas abas de dados"
    for w in data_sheet_widths_px:
        assert w == pytest.approx(ewb.IMAGE_TARGET_WIDTH_PX, abs=2), \
            f"largura exibida {w:.1f}px longe do alvo {ewb.IMAGE_TARGET_WIDTH_PX}px"

    assert resumo_boxes, "nenhuma imagem encontrada no Resumo gráficos"
    for _, _, cx in resumo_boxes:
        w = cx / EMU_PER_PX
        assert w == pytest.approx(ewb.RESUMO_IMAGE_TARGET_WIDTH_PX, abs=2), \
            f"largura exibida {w:.1f}px longe do alvo {ewb.RESUMO_IMAGE_TARGET_WIDTH_PX}px"

    resumo_boxes.sort(key=lambda b: b[0])
    for (y1, cy1, _), (y2, _, _) in zip(resumo_boxes, resumo_boxes[1:]):
        assert y1 + cy1 <= y2, "figuras do Resumo gráficos se sobrepõem"


def test_workbook_include_images_false_drops_summary_and_media(export_request):
    """Variante "somente tabelas" (aba EXPORT, item opcional do pedido):
    mesmas abas de dados, mesma formatacao, mas sem Resumo gráficos e sem
    nenhuma imagem embutida -- arquivo mais leve pra quem so quer os
    numeros."""
    import openpyxl
    xlsx_bytes = ewb.build_workbook(export_request, include_images=False)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

    assert "Resumo gráficos" not in wb.sheetnames
    expected = [
        "Inputs",
        "Design 297 K", "Design 339 K", "Design 422 K",
        "Sim 5.00 ft", "Sim 10.00 ft", "Sim 15.00 ft", "Sim 20.00 ft",
        "Skin 0.8 bbl-min", "Skin 1.6 bbl-min", "Skin 3.2 bbl-min",
    ]
    assert wb.sheetnames == expected

    z = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    media = [n for n in z.namelist() if n.startswith("xl/media/")]
    assert media == []


FIGURE_TITLE_FILL_ARGB = "FF1F4E78"


def _assert_title_bar_style(cell):
    """MESMO estilo do cabecalho 'INPUT PARAMETERS' da aba Inputs: fundo
    solido #1F4E78, fonte branca em negrito 12pt, centralizado H+V."""
    assert cell.fill.fill_type == "solid"
    assert cell.fill.fgColor.rgb == FIGURE_TITLE_FILL_ARGB
    assert cell.font.bold is True
    assert cell.font.size == 12
    assert cell.font.color.rgb == "FFFFFFFF"
    assert cell.alignment.horizontal == "center"
    assert cell.alignment.vertical == "center"


def test_workbook_figure_titles_replace_filename_with_formatted_bar(export_request):
    """Bug real: a barra de titulo das figuras mostrava o nome interno do
    arquivo ("simulation_all") como texto solto, sem formatacao. Cada aba
    com imagem agora tem um titulo legivel e simples -- so o nome do
    grafico (Simulation Chart / Design Plot / Skin Evolution), sem sufixo
    de alvo/temperatura/vazao -- na MESMA celula mesclada estilo INPUT
    PARAMETERS, na linha imediatamente acima da imagem -- nunca o nome
    interno do PNG."""
    import openpyxl
    xlsx_bytes = ewb.build_workbook(export_request)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

    per_sheet_titles = {
        "Sim 5.00 ft": "Simulation Chart",
        "Design 297 K": "Design Plot",
        "Skin 0.8 bbl-min": "Skin Evolution",
    }
    for sheet_name, expected_text in per_sheet_titles.items():
        ws = wb[sheet_name]
        assert ws._images, f"{sheet_name} sem imagem"
        img = ws._images[0]
        img_row0, img_col0 = img.anchor._from.row, img.anchor._from.col

        # titulo na linha IMEDIATAMENTE acima da imagem (0-indexed img_row0
        # == 1-indexed excel row do titulo); imagem ancorada na linha seguinte.
        title_row, title_col = img_row0, img_col0 + 1
        cell = ws.cell(row=title_row, column=title_col)
        assert cell.value == expected_text
        _assert_title_bar_style(cell)
        assert ws.row_dimensions[title_row].height == 24
        assert any(cell.coordinate in rng for rng in ws.merged_cells.ranges), \
            f"{sheet_name}: titulo nao esta numa celula mesclada"

    ws_resumo = wb["Resumo gráficos"]
    expected_resumo_titles = {
        "Simulation Chart",
        "Design Plot",
        "Skin Evolution",
    }
    seen_titles = set()
    for img in ws_resumo._images:
        img_row0, img_col0 = img.anchor._from.row, img.anchor._from.col
        title_row, title_col = img_row0, img_col0 + 1
        cell = ws_resumo.cell(row=title_row, column=title_col)
        assert cell.value in expected_resumo_titles, f"titulo inesperado: {cell.value!r}"
        seen_titles.add(cell.value)
        _assert_title_bar_style(cell)
        assert ws_resumo.row_dimensions[title_row].height == 24
        assert any(cell.coordinate in rng for rng in ws_resumo.merged_cells.ranges)
    assert seen_titles == expected_resumo_titles


def test_workbook_design_plot_multi_target_highlights(export_request):
    import openpyxl
    xlsx_bytes = ewb.build_workbook(export_request)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

    ws297 = wb["Design 297 K"]
    notes_297 = [c.value for c in ws297["F"][1:] if c.value]
    assert len(notes_297) == 4  # todos os 4 alvos atingidos, um por linha

    ws339 = wb["Design 339 K"]
    notes_339 = [c.value for c in ws339["F"][1:] if c.value]
    assert any("15" in n and "20" in n and "não atingidos" in n for n in notes_339)

    ws422 = wb["Design 422 K"]
    notes_422 = [c.value for c in ws422["F"][1:] if c.value]
    assert any("10, 15 e 20" in n for n in notes_422)


def test_workbook_values_match_source_curves(export_request):
    """As celulas da aba Sim tem que bater com o payload -- nenhum
    recalculo no meio do caminho."""
    import openpyxl
    xlsx_bytes = ewb.build_workbook(export_request)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

    curve = next(c for c in export_request.curves if c.target_label == "5.00 ft")
    ws = wb["Sim 5.00 ft"]
    for i in range(len(curve.flowratepoints)):
        row = i + 2
        assert ws.cell(row=row, column=1).value == pytest.approx(curve.flowratepoints[i])
        assert ws.cell(row=row, column=2).value == pytest.approx(curve.acidvolumepoints[i])


# --- Figures zip ---------------------------------------------------------

EXPECTED_STEMS = {
    "simulation_5.00ft", "simulation_10.00ft", "simulation_15.00ft", "simulation_20.00ft", "simulation_all",
    "design_297.04K", "design_338.71K", "design_422.04K", "design_all",
    "skin_0.8bblmin", "skin_1.6bblmin", "skin_3.2bblmin", "skin_all",
}


@pytest.mark.parametrize("size,expected_mm", [("single", (90.0, 70.0)), ("double", (190.0, 120.0))])
def test_figures_zip_contents_dpi_and_size(export_request, size, expected_mm):
    figures = ewb.generate_all_figures(export_request, size=size)
    assert set(figures.keys()) == EXPECTED_STEMS

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for stem, png in figures.items():
            zf.writestr(f"{stem}.png", png)
    buf.seek(0)
    z = zipfile.ZipFile(buf)
    names = {n[:-4] for n in z.namelist()}
    assert names == EXPECTED_STEMS

    w_mm, h_mm = expected_mm
    for name in z.namelist():
        im = Image.open(io.BytesIO(z.read(name)))
        dpi = im.info.get("dpi", (72, 72))
        assert dpi[0] == pytest.approx(300, abs=1) and dpi[1] == pytest.approx(300, abs=1), name

        implied_w = im.size[0] / dpi[0] * MM_PER_INCH
        implied_h = im.size[1] / dpi[1] * MM_PER_INCH
        assert implied_w <= w_mm * (1 + 1e-6), f"{name}: largura {implied_w:.1f}mm > {w_mm}mm"
        assert implied_h <= h_mm * (1 + 1e-6), f"{name}: altura {implied_h:.1f}mm > {h_mm}mm"
        assert implied_w >= w_mm * (1 - SIZE_TOLERANCE), f"{name}: largura {implied_w:.1f}mm bem menor que {w_mm}mm"


def test_font_falls_back_safely():
    assert ep.FONT_FAMILY  # nunca vazio -- sempre resolve pra algo (DejaVu Sans no pior caso)


# --- Flowing Fraction (f) -------------------------------------------------
# Bug real (2026-09-18): a celula "Flowing Fraction (f)" na aba Inputs
# sempre mostrava "nao disponivel", pra QUALQUER rock_type -- self.f
# (RadialCurveMaster.__init__, PVBTradialFunc.py) era calculado e usado em
# _integral_fechada, mas get_adjusted_parameters() nunca o devolvia, e o
# frontend lia de curve.metadata (janela de validade, RadialCurveValidity),
# que nunca teve essa chave. Um teste so de "o campo existe" passaria com
# f=1.0 chumbado (o default de ROCKFLOWFRACTION.get); por isso comparamos
# DOIS rocks com fracoes diferentes e conferimos que os valores batem com a
# tabela (PVBTfunc.ROCKFLOWFRACTION) e sao DIFERENTES entre si.
from app.services.PVBTfunc import ROCKFLOWFRACTION

FF_ROCKS = [("Indiana Limestone", 1.00), ("Edwards White", 0.52)]


def _flowing_fraction_cell(rock_type, expected_table_value):
    assert ROCKFLOWFRACTION[rock_type] == pytest.approx(expected_table_value)

    data = RadialCurveInput(
        simulation_id="TEST-FF", system=_system(), geometry=_geometry(),
        radial_targets=RadialTargetsInput(target_mode="length", targets=TARGETS),
        flowrate_sweep=FlowrateSweepInput(min=SWEEP_MIN_BBL, max=SWEEP_MAX_BBL, steps=STEPS),
    )
    data.system.rock_type = rock_type
    output = calculate_pvbt_radial(data)

    # f chega no cliente por parameters.f (get_adjusted_parameters), NAO por
    # curve.metadata (esse e so a janela de validade da vazao).
    assert "f" in output["parameters"], "get_adjusted_parameters() nao devolve mais 'f'"
    f_value = output["parameters"]["f"]
    assert f_value == pytest.approx(expected_table_value)

    curves = [RadialCurveResult(**c) if not isinstance(c, RadialCurveResult) else c
              for c in output["curves"]]
    req = RadialExportRequest(
        inputs=ExportInputs(
            simulation_id="TEST-FF", rock_type=rock_type, acid_type=ACID,
            acid_concentration=CONC, porosity=POROSITY, temperature_k=TEMP_K,
            wellbore_size_in=3.0, wellbore_mode="diameter",
            wellbore_radius_in=WELLBORE_RADIUS_IN, payzone_thickness_ft=PAYZONE_FT,
            flowrate_min_bbl_min=SWEEP_MIN_BBL, flowrate_max_bbl_min=SWEEP_MAX_BBL,
            number_of_steps=STEPS, targets_label=", ".join(c.target_label for c in curves),
            flowing_fraction=f_value,
        ),
        curves=curves, design_series=[], skin_series={},
        options=RadialExportOptions(target_lengths=TARGETS, show_optimum_path=True, show_validity_band=True),
    )

    import openpyxl
    xlsx_bytes = ewb.build_workbook(req, include_images=False)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Inputs"]
    cell_value = next(row[1].value for row in ws.iter_rows(min_row=2, max_col=2)
                       if row[0].value == "Flowing Fraction (f)")
    assert cell_value == pytest.approx(expected_table_value)
    return cell_value


@pytest.mark.parametrize("rock_type,expected", FF_ROCKS)
def test_flowing_fraction_matches_rock_table(rock_type, expected):
    _flowing_fraction_cell(rock_type, expected)


def test_flowing_fraction_differs_between_rocks():
    values = [_flowing_fraction_cell(rock, expected) for rock, expected in FF_ROCKS]
    assert values[0] != values[1], \
        "Flowing Fraction (f) saiu igual para rocks com fracoes diferentes na tabela -- suspeita de valor chumbado"


# --- Baseline visual (pytest-mpl) ---------------------------------------

def _baseline_curves():
    return [
        ep.SimCurveData(
            label="5.00 ft", color=ep.color_for_index(0),
            flowratepoints=[0.1, 0.5, 1.0, 2.0, 4.0],
            acidvolumepoints=[80.0, 25.0, 22.0, 28.0, 45.0],
            q_opt=1.0, validity_min=0.1, validity_max=4.0,
        ),
    ]


@pytest.mark.mpl_image_compare(baseline_dir="baseline", filename="simulation_baseline.png", tolerance=5)
def test_simulation_figure_matches_baseline():
    return ep.build_simulation_figure(_baseline_curves(), size="single")
