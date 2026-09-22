"""
tests/test_export_linear.py

Export linear (.xlsx com figura matplotlib + otimo). Os dados vem da MESMA
funcao que a rota /pvbtcurve executa (calculate_pvbt) -- ou seja, q_opt aqui
e o que a tela mostra ("opt ..." em SimuCard), nunca recalculado pelo export.

    venv/Scripts/python.exe -m pytest tests/test_export_linear.py -v

Definir PVBT_EXPORT_DUMP_DIR=<pasta> grava os .xlsx/.png gerados (verificacao
visual manual).
"""
import io
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import numpy as np
import openpyxl
import pytest
from PIL import Image

from app.routes.pvbtCurve import calculate_pvbt
from app.schemas import (
    LinearExperimentalCurve,
    LinearExportRequest,
    LinearModelCurve,
    PVBtInputCurve,
)
from app.services import export_workbook_linear as ewl

ROCK = "Indiana Limestone"
ACID = "HCl With Inhibitor Corrosion"
DUMP = os.environ.get("PVBT_EXPORT_DUMP_DIR")


def run_curve(curve_id, temp_c, qmin, qmax, steps=50) -> LinearModelCurve:
    """Chama a rota real e monta o payload como o front (state.resultCurves)."""
    inp = PVBtInputCurve(
        acid_type=ACID, acid_concentration=0.15, core_diameter=1.5, core_length=6.0,
        core_porosity=0.15, rock_type=ROCK, temperature=temp_c,
        flowrate=qmax, minimum_flowrate=qmin, step_numbers=steps,
    )
    out = calculate_pvbt(inp)
    return LinearModelCurve(
        id=curve_id, rock_type=ROCK, acid_type=ACID, acid_concentration=0.15,
        porosity=0.15, temperature_c=temp_c, core_length_in=6.0, core_diameter_in=1.5,
        flowratepoints=out["flowratepoints"], pvbtpoints=out["pvbtpoints"],
        insterticialvelocity=out["insterticialvelocity"], ida=out["ida"],
        wormholevelocity=out["wormholevelocity"], volumetobt=out["volumetobt"],
        timetobt=out["timetobt"], darcyvelocity=out["darcyvelocity"],
        within_validity_range=out["within_validity_range"], metadata=out["metadata"],
    )


def experimental(curve_id="Exp: core A"):
    return LinearExperimentalCurve(
        id=curve_id, flowratepoints=[0.5, 1.0, 2.0, 5.0, 12.0],
        pvbtpoints=[3.1, 1.9, 1.5, 1.7, 2.8],
    )


def build(req, include_images=True) -> openpyxl.Workbook:
    data = ewl.build_linear_workbook(req, include_images=include_images)
    if DUMP:
        os.makedirs(DUMP, exist_ok=True)
        tag = "_".join(c.id.replace(" ", "") for c in req.curves) or "exp"
        suffix = "" if include_images else "_tables"
        with open(os.path.join(DUMP, f"linear_{tag}{suffix}.xlsx"), "wb") as f:
            f.write(data)
    return openpyxl.load_workbook(io.BytesIO(data))


def summary_values(ws) -> dict:
    return {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, 8)}


def data_rows(ws):
    header_row = next(r for r in range(1, 20) if str(ws.cell(r, 1).value or "").startswith("q0"))
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        rows.append([ws.cell(r, c).value for c in range(1, 10)])
    return header_row, rows


@pytest.fixture(scope="module")
def two_curves():
    return [run_curve("T25", 25.0, 0.1, 30.0), run_curve("T50", 50.0, 0.1, 30.0)]


def test_metadata_has_exact_pvbt_at_q_opt(two_curves):
    for c in two_curves:
        assert c.metadata is not None
        grid_min = min(v for v in c.pvbtpoints if v is not None)
        assert c.metadata["pvbt_at_q_opt"] <= grid_min * (1 + 1e-9)


def test_exported_q_opt_equals_source_value(two_curves):
    wb = build(LinearExportRequest(curves=two_curves, experimental_curves=[experimental()]))
    for c in two_curves:
        ws = wb[f"Sim {c.id}"]
        s = summary_values(ws)
        q_opt = c.metadata["q_opt_cm3_min"]
        assert s["q_opt [cm³/min]"] == pytest.approx(q_opt, rel=1e-14)
        assert s["PVBT at q_opt"] == pytest.approx(c.metadata["pvbt_at_q_opt"], rel=1e-14)
        assert s["Recommended window min = q_opt/10 [cm³/min]"] == pytest.approx(q_opt / 10)
        assert s["Recommended window max = 10·q_opt [cm³/min]"] == pytest.approx(q_opt * 10)
        _, rows = data_rows(ws)
        notes = [r[8] for r in rows if r[8]]
        assert len(notes) == 1
        assert f"q_opt = {ewl.fmt_flow(q_opt)} cm³/min" in notes[0]
        assert notes[0].startswith("Minimum PVBT of this simulation")


def test_two_optima_are_kept_distinct(two_curves):
    wb = build(LinearExportRequest(curves=two_curves))
    ws = wb["Sim T25"]
    s = summary_values(ws)
    assert "Lowest PVBT in sweep (grid-dependent)" in s
    assert s["Lowest PVBT in sweep (grid-dependent)"] >= s["PVBT at q_opt"]
    _, rows = data_rows(ws)
    noted = next(r for r in rows if r[8])
    assert noted[1] == pytest.approx(s["Lowest PVBT in sweep (grid-dependent)"], rel=1e-14)
    assert noted[0] == pytest.approx(s["Flowrate at lowest swept PVBT [cm³/min]"], rel=1e-14)


def test_edge_of_sweep_note_gives_true_q_opt():
    c = run_curve("edge", 24.05, 5.0, 10.0, steps=20)
    q_opt = c.metadata["q_opt_cm3_min"]
    assert q_opt == pytest.approx(1.7054, abs=1e-3)
    wb = build(LinearExportRequest(curves=[c]))
    _, rows = data_rows(wb["Sim edge"])
    noted = [(i, r[8]) for i, r in enumerate(rows) if r[8]]
    assert len(noted) == 1
    idx, note = noted[0]
    assert idx in (0, len(rows) - 1)
    assert note == f"Minimum at the edge of the simulated range; q_opt = {ewl.fmt_flow(q_opt)} cm³/min"
    assert "Minimum PVBT of this simulation" not in note
    fig = ewl.build_linear_figure(LinearExportRequest(curves=[c]))
    ax = fig.axes[0]
    assert not [ln for ln in ax.lines if ln.get_marker() == "o" and list(ln.get_xdata())]


def test_no_metadata_degrades_without_q_opt():
    c = run_curve("nometa", 25.0, 0.1, 30.0)
    c = c.model_copy(update={"metadata": None})
    wb = build(LinearExportRequest(curves=[c]))
    ws = wb["Sim nometa"]
    s = summary_values(ws)
    assert s["q_opt [cm³/min]"] == "not available"
    _, rows = data_rows(ws)
    note = next(r[8] for r in rows if r[8])
    assert "q_opt" not in note


def test_experimental_sheet_present_and_bare(two_curves):
    exp = experimental()
    wb = build(LinearExportRequest(curves=two_curves, experimental_curves=[exp]))
    assert "Experimental" in wb.sheetnames
    ws = wb["Experimental"]
    assert [ws.cell(1, c).value for c in (1, 2, 3)] == ["Curve ID", "q0 [cm³/min]", "PVBt"]
    got = [(ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(2, 7)]
    assert got == [(exp.id, q, y) for q, y in zip(exp.flowratepoints, exp.pvbtpoints)]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Note" not in text and "OPTIMUM SUMMARY" not in text
    assert "Sim Exp" not in " ".join(wb.sheetnames)


def test_figure_nonblank_log_log_markers_optimum_band(two_curves):
    req = LinearExportRequest(curves=two_curves, experimental_curves=[experimental()])
    wb = build(req)
    ws = wb["Figures"]
    assert len(ws._images) == 1
    png = ws._images[0]._data()
    im = Image.open(io.BytesIO(png)).convert("L")
    arr = np.asarray(im)
    assert arr.shape[0] > 300 and arr.shape[1] > 300
    assert (arr < 200).mean() > 0.01
    if DUMP:
        im.save(os.path.join(DUMP, "linear_figure.png"))

    ax = ewl.build_linear_figure(req).axes[0]
    assert ax.get_xscale() == "log" and ax.get_yscale() == "log"
    assert ax.get_xlabel() == "Flowrate, cm³/min" and ax.get_ylabel() == "PVBt"
    exp_lines = [ln for ln in ax.lines if ln.get_label() == "Exp: core A"]
    assert len(exp_lines) == 1
    assert exp_lines[0].get_linestyle() == "None" and exp_lines[0].get_marker() == "s"
    for c in two_curves:
        qo, yo = c.metadata["q_opt_cm3_min"], c.metadata["pvbt_at_q_opt"]
        assert any(
            ln.get_marker() == "o" and list(ln.get_xdata()) == [qo] and list(ln.get_ydata()) == [yo]
            for ln in ax.lines
        ), c.id
    assert len(ax.patches) >= 2
    assert not [ln for ln in ax.lines if ln.get_label() == "Exp: core A" and ln.get_marker() == "o"]


def test_tables_only_variant_has_optimum_but_no_figure(two_curves):
    wb = build(LinearExportRequest(curves=two_curves, experimental_curves=[experimental()]),
               include_images=False)
    assert "Figures" not in wb.sheetnames
    assert summary_values(wb["Sim T25"])["q_opt [cm³/min]"] == pytest.approx(
        two_curves[0].metadata["q_opt_cm3_min"], rel=1e-14)


def test_inputs_sheet_units_and_columns(two_curves):
    wb = build(LinearExportRequest(curves=two_curves))
    ws = wb["Inputs"]
    labels = {ws.cell(r, 1).value: [ws.cell(r, 2).value, ws.cell(r, 3).value] for r in range(2, 20)
              if ws.cell(r, 1).value}
    assert labels["Simulation ID"] == ["T25", "T50"]
    assert labels["Flow Regime"] == ["linear", "linear"]
    assert labels["Temperature (°C)"] == [25.0, 50.0]
    assert labels["Number of steps"] == [50, 50]
    assert labels["Flowrate Sweep Min (cm³/min)"][0] == pytest.approx(0.1)
    assert labels["Flowrate Sweep Max (cm³/min)"][0] == pytest.approx(30.0)
    _, rows = data_rows(wb["Sim T25"])
    header = [wb["Sim T25"].cell(r, 1).value for r in range(1, 12)]
    assert any(h == "q0 [cm³/min]" for h in header)
